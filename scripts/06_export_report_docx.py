"""Convert docs/PROJECT_REPORT.md into a Word document for submission."""
from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))


def add_md_line(doc: Document, line: str) -> None:
    raw = line.rstrip()
    if not raw:
        return
    if raw.startswith("# "):
        doc.add_heading(raw[2:].strip(), level=0)
        return
    if raw.startswith("## "):
        doc.add_heading(raw[3:].strip(), level=1)
        return
    if raw.startswith("### "):
        doc.add_heading(raw[4:].strip(), level=2)
        return
    if raw.startswith("---"):
        return
    if raw.startswith("|") and "---" not in raw:
        # Collect tables elsewhere; skip single rows here
        return
    if raw.startswith("```"):
        return
    if raw.startswith("> "):
        p = doc.add_paragraph(raw[2:].strip())
        p.runs[0].italic = True if p.runs else None
        return
    if raw.startswith("- "):
        doc.add_paragraph(raw[2:].strip(), style="List Bullet")
        return
    if re.match(r"^\d+\.\s", raw):
        doc.add_paragraph(re.sub(r"^\d+\.\s+", "", raw), style="List Number")
        return
    # Inline code / bold lite cleanup
    text = re.sub(r"`([^`]+)`", r"\1", raw)
    text = re.sub(r"\*\*([^*]+)\*\*", r"\1", text)
    doc.add_paragraph(text)


def add_table_from_rows(doc: Document, rows: list[list[str]]) -> None:
    if len(rows) < 2:
        return
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    for i, row in enumerate(rows):
        for j, cell in enumerate(row):
            table.rows[i].cells[j].text = cell.strip()


def parse_markdown(md: str, doc: Document) -> None:
    lines = md.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        if line.strip().startswith("|"):
            rows: list[list[str]] = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                row = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(set(c) <= set("-: ") for c in row):
                    rows.append(row)
                i += 1
            add_table_from_rows(doc, rows)
            continue
        if line.strip().startswith("```"):
            i += 1
            block: list[str] = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                block.append(lines[i])
                i += 1
            p = doc.add_paragraph("\n".join(block))
            for run in p.runs:
                run.font.name = "Consolas"
                run.font.size = Pt(9)
            i += 1
            continue
        add_md_line(doc, line)
        i += 1


def main() -> None:
    md_path = ROOT / "docs" / "PROJECT_REPORT.md"
    out_path = ROOT / "docs" / "Aegis_IDS_Project_Report.docx"
    figures = ROOT / "models" / "trained_models" / "figures"

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    title = doc.add_heading("Aegis IDS", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub = doc.add_paragraph(
        "Intelligent Machine Learning–Based Intrusion Detection "
        "and Attack–Defense Simulation System"
    )
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER

    parse_markdown(md_path.read_text(encoding="utf-8"), doc)

    doc.add_heading("Appendix C — Evaluation figures", level=1)
    for name, caption in [
        ("binary_confusion_matrix.png", "Figure C1. Binary confusion matrix (IID test sample)"),
        ("binary_roc.png", "Figure C2. Binary ROC curve (IID)"),
        ("multiclass_confusion_matrix.png", "Figure C3. Multiclass confusion matrix (IID, six attack classes)"),
        ("feature_importance.png", "Figure C4. Top feature importances"),
        ("temporal_binary_iid_vs_holdout.png", "Figure C5. Binary metrics: IID vs Friday vs Mon–Thu samples"),
        ("temporal_multiclass_iid_vs_holdout.png", "Figure C6. Multiclass: IID (6-class) vs Friday subset"),
        ("temporal_friday_multiclass_confusion.png", "Figure C7. Friday temporal multiclass confusion matrix (present families)"),
        ("temporal_friday_classwise.png", "Figure C8. Friday temporal class-wise support vs F1"),
        ("temporal_generalization_summary.png", "Figure C9. Temporal generalization summary card"),
        ("drift_psi_train_vs_test.png", "Figure C10. Feature PSI train→test (IID drift check)"),
        ("drift_label_distribution.png", "Figure C11. Label distribution train vs test (IID)"),
        ("error_binary_fp_fn_counts.png", "Figure C12. Binary FP/FN residual counts (IID test)"),
        ("error_multiclass_top_confusions.png", "Figure C13. Top multiclass residual confusions"),
        ("drift_error_research_summary.png", "Figure C14. Drift + error research summary card"),
    ]:
        path = figures / name
        if path.exists():
            doc.add_picture(str(path), width=Inches(5.5))
            cap = doc.add_paragraph(caption)
            cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if cap.runs:
                cap.runs[0].italic = True
        else:
            miss = doc.add_paragraph(f"[Missing figure: {name}]")
            if miss.runs:
                miss.runs[0].italic = True

    doc.save(out_path)
    print(f"Wrote {out_path}")


if __name__ == "__main__":
    main()
