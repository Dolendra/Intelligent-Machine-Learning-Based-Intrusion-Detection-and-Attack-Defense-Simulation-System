"""Generate final viva PowerPoint (18 slides) from PROJECT_REPORT.md numbers.

Aligned to tags v1.1-research (frozen ML) and v2.0-aegis-productionized.

  python scripts/08_export_presentation_pptx.py
  → docs/Aegis_IDS_Viva_Presentation.pptx
"""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
FIGS = ROOT / "models" / "trained_models" / "figures"
OUT = ROOT / "docs" / "Aegis_IDS_Viva_Presentation.pptx"

BG = RGBColor(0x07, 0x11, 0x1A)
ACCENT = RGBColor(0x3E, 0xC7, 0xC2)
TEXT = RGBColor(0xE7, 0xF0, 0xF5)
MUTED = RGBColor(0x8F, 0xA6, 0xB5)
WARN = RGBColor(0xF0, 0xB4, 0x29)


def _set_run(run, size=20, bold=False, color=TEXT):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = "Calibri"


def add_blank(prs: Presentation):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), prs.slide_width, prs.slide_height
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = BG
    shape.line.fill.background()
    spTree = slide.shapes._spTree
    sp = shape._element
    spTree.remove(sp)
    spTree.insert(2, sp)
    return slide


def title_block(slide, title: str, subtitle: str | None = None):
    box = slide.shapes.add_textbox(Inches(0.6), Inches(0.35), Inches(12.2), Inches(1.0))
    tf = box.text_frame
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = title
    _set_run(run, 30, True, ACCENT)
    if subtitle:
        box2 = slide.shapes.add_textbox(Inches(0.6), Inches(1.05), Inches(12.2), Inches(0.5))
        p2 = box2.text_frame.paragraphs[0]
        r2 = p2.add_run()
        r2.text = subtitle
        _set_run(r2, 14, False, MUTED)


def bullets(slide, lines: list[str], top=1.7, left=0.7, width=12, size=18):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(5.4))
    tf = box.text_frame
    tf.word_wrap = True
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.level = 0
        run = p.add_run()
        run.text = "•  " + line
        _set_run(run, size, False, TEXT)
        p.space_after = Pt(7)


def footer(slide, text: str = "Aegis IDS  ·  v2.0-aegis-productionized  ·  Research baseline v1.1-research"):
    box = slide.shapes.add_textbox(Inches(0.6), Inches(7.05), Inches(12.2), Inches(0.35))
    p = box.text_frame.paragraphs[0]
    r = p.add_run()
    r.text = text
    _set_run(r, 11, False, MUTED)


def add_pic(slide, name: str, left: float, top: float, height: float | None = None, width: float | None = None):
    path = FIGS / name
    if not path.exists():
        return False
    kwargs = {}
    if height is not None:
        kwargs["height"] = Inches(height)
    if width is not None:
        kwargs["width"] = Inches(width)
    slide.shapes.add_picture(str(path), Inches(left), Inches(top), **kwargs)
    return True


def main() -> None:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # 1 Title
    s = add_blank(prs)
    box = s.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(11.7), Inches(4.5))
    tf = box.text_frame
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = "Aegis IDS"
    _set_run(r, 48, True, ACCENT)
    p2 = tf.add_paragraph()
    p2.alignment = PP_ALIGN.CENTER
    r2 = p2.add_run()
    r2.text = "Intelligent ML-Based Intrusion Detection\n& Attack–Defense Simulation System"
    _set_run(r2, 24, False, TEXT)
    p3 = tf.add_paragraph()
    p3.alignment = PP_ALIGN.CENTER
    r3 = p3.add_run()
    r3.text = (
        "\nv1.1-research  →  v2.0-aegis-productionized\n"
        "Decision Tree @ 0.85  ·  Random Forest (attack-only)  ·  P12 PASS"
    )
    _set_run(r3, 15, False, MUTED)
    p4 = tf.add_paragraph()
    p4.alignment = PP_ALIGN.CENTER
    r4 = p4.add_run()
    r4.text = "\nTeam members  ·  College  ·  Year"
    _set_run(r4, 14, False, MUTED)

    # 2 Problem
    s = add_blank(prs)
    title_block(s, "Problem", "Why a classic IDS alert is not enough")
    bullets(
        s,
        [
            "Typical IDS output: “Attack detected.”",
            "Analysts still need: attack type · why · severity · recommended action",
            "Evaluators need a clear attack → defense story with honest limits",
            "Gap: detection alone ≠ security decision support",
        ],
    )
    footer(s)

    # 3 Motivation
    s = add_blank(prs)
    title_block(s, "Motivation & objectives")
    bullets(
        s,
        [
            "Bridge detection → explainable decision support with human control",
            "Detect & classify CICIDS2017 flows with a frozen research baseline",
            "Explain (SHAP), score risk, recommend advisory defenses",
            "Simulate the lifecycle; apply CONTROLLED responses with approval",
            "Productionize safely (P0–P12) without inflating research metrics",
        ],
        size=17,
    )
    footer(s)

    # 4 Existing limitations
    s = add_blank(prs)
    title_block(s, "Limitations of many existing demos / systems")
    bullets(
        s,
        [
            "High CICIDS accuracy without temporal / drift honesty",
            "Opaque models → weak analyst trust",
            "Recommendations missing or auto-enforcement without safeguards",
            "Simulation efficacy confused with measured mitigation",
            "Weak story for auth, audit, persistence, and fail-safe recovery",
        ],
        size=17,
    )
    footer(s)

    # 5 Proposed
    s = add_blank(prs)
    title_block(s, "Proposed: Aegis IDS", "One coherent pipeline")
    bullets(
        s,
        [
            "Detect → Classify → Explain → Risk → Recommend → Approve → CONTROLLED response → Simulate",
            "Contribution: integration + honest evaluation — not a new IDS algorithm claim",
            "Hard boundary: no live firewall/EDR · no unapproved automatic response",
            "Tags: v1.1-research (freeze) · v2.0-aegis-productionized (prototype)",
        ],
        size=17,
    )
    footer(s)

    # 6 Architecture
    s = add_blank(prs)
    title_block(s, "Architecture")
    bullets(
        s,
        [
            "React SOC UI  ↔  FastAPI  ↔  ML + Risk + XAI + Response + Simulation",
            "Adapters: DRY_RUN | CONTROLLED (TestNetworkAdapter) | LIVE (forbidden)",
            "Persistence: SQLite incidents / actions / append-only audit",
            "Observability: /api/ready · ops metrics · System page",
        ],
        left=0.6,
        width=7.0,
        size=16,
    )
    arch = ROOT / "docs" / "architecture.png"
    if arch.exists():
        s.shapes.add_picture(str(arch), Inches(7.5), Inches(1.7), width=Inches(5.3))
    footer(s)

    # 7 Dataset
    s = add_blank(prs)
    title_block(s, "Dataset & preprocessing", "CICIDS2017 MachineLearningCVE")
    bullets(
        s,
        [
            "Clean · normalize labels · drop rare classes (<50 samples)",
            "~2.52M flows after cleaning; stratified 70% / 10% / 20%",
            "Splits: 1,764,525 / 252,075 / 504,151 · random_state=42",
            "StandardScaler + SelectKBest(f_classif, k=40) · dual selectors · train-only fit",
            "Families: Bot · BruteForce · DDoS · DoS · PortScan · WebAttack",
        ],
        size=16,
    )
    footer(s)

    # 8 ML architecture
    s = add_blank(prs)
    title_block(s, "ML architecture (two-stage)")
    bullets(
        s,
        [
            "Stage 1 — Binary: BENIGN vs ATTACK → Decision Tree @ threshold 0.85",
            "Stage 2 — Attack-only multiclass → Random Forest (no BENIGN class)",
            "Candidates compared: Logistic Regression, DT, RF, XGBoost",
            "Same feature pipeline for fair algorithm ablation",
        ],
        size=18,
    )
    footer(s)

    # 9 Results
    s = add_blank(prs)
    title_block(s, "Model results (research IID)", "Frozen artifacts — not a live-network guarantee")
    bullets(
        s,
        [
            "Binary DT @ 0.85: F1 0.99048 · Recall 0.997 · PR-AUC 0.99744 · ROC-AUC 0.99916",
            "Binary residuals: FP 1,377 · FN 255 (FPR 0.00329 / FNR 0.00300)",
            "Multiclass RF: macro-F1 0.99813 · weighted-F1 0.99979 · accuracy 0.99979",
            "Cite as: performance on the stratified IID holdout",
        ],
        left=0.55,
        width=6.4,
        size=15,
    )
    add_pic(s, "binary_confusion_matrix.png", 7.1, 1.65, height=2.55)
    add_pic(s, "binary_roc.png", 7.1, 4.3, height=2.4)
    footer(s)

    # 10 Why DT
    s = add_blank(prs)
    title_block(s, "Why Decision Tree (not XGBoost)?", "Multi-objective selection — recall first")
    bullets(
        s,
        [
            "Weights: recall 0.30 · F1 0.25 · PR-AUC 0.20 · FPR 0.15 · latency 0.10",
            "XGBoost can win raw F1; Decision Tree wins attack recall → selected",
            "Multiclass: Random Forest edges XGBoost on macro/weighted blend",
            "For IDS, missed attacks (FN) are costly — selection policy matters",
        ],
        left=0.55,
        width=6.5,
        size=16,
    )
    add_pic(s, "model_binary_f1_vs_recall.png", 7.2, 1.85, width=5.6)
    footer(s)

    # 11 XAI + Risk
    s = add_blank(prs)
    title_block(s, "Explainability & risk", "Decision support — not causality")
    bullets(
        s,
        [
            "SHAP primary · LIME secondary (Detection → Why? tabs)",
            "Risk blend: 50% attack base · 25% confidence · 15% intensity · 10% asset",
            "Severity bands: LOW / MEDIUM / HIGH / CRITICAL (project conventions)",
            "Recommendations are advisory until human approval",
        ],
        left=0.55,
        width=6.5,
        size=16,
    )
    add_pic(s, "feature_importance.png", 7.2, 1.85, width=5.5)
    footer(s)

    # 12 Simulation
    s = add_blank(prs)
    title_block(s, "Attack–defense simulation", "Visualization only — assumptions ≠ measurements")
    bullets(
        s,
        [
            "idle → normal → attack_start → attack_impact → detected → recommended → defended → recovered",
            "Topology: Attacker → Firewall → Router → Server / PCs + IDS",
            "Simulation priors (examples): DDoS 0.82 · DoS 0.78 (config.yaml)",
            "Safety: controlled visualization — no real cyberattacks / no live mitigation",
        ],
        size=16,
    )
    footer(s)

    # 13 Controlled response
    s = add_blank(prs)
    title_block(s, "Controlled response architecture", "Human-in-the-loop safety")
    bullets(
        s,
        [
            "Lifecycle: propose → dry-run → approve → execute → verify → rollback",
            "Modes: DRY_RUN · CONTROLLED (TestNetworkAdapter) · LIVE forbidden",
            "RBAC: viewer (read) · analyst (propose) · responder/admin (approve/rollback)",
            "Verify failure → rollback · reclaim never auto-replays dangerous actions",
        ],
        size=16,
    )
    footer(s)

    # 14 P11
    s = add_blank(prs)
    title_block(s, "P11 — Empirical mitigation", "Controlled lab measurements vs simulation priors")
    bullets(
        s,
        [
            "Conditions: no_defense vs recommendation_only vs controlled_response (N=10)",
            "DDoS BLOCK_SOURCE: measured effectiveness 1.00  (sim prior 0.82 · Δ +0.18)",
            "DoS RATE_LIMIT: measured effectiveness 0.80  (sim prior 0.78 · Δ +0.02)",
            "Recommendation-only does not change traffic — detects ≠ mitigates",
            "CONTROLLED lab only — not production firewall/EDR effectiveness",
        ],
        size=16,
    )
    # callout
    box = s.shapes.add_textbox(Inches(0.7), Inches(6.15), Inches(12), Inches(0.55))
    p = box.text_frame.paragraphs[0]
    r = p.add_run()
    r.text = "Keep both numbers: simulation assumptions stay labelled assumptions."
    _set_run(r, 14, True, WARN)
    footer(s)

    # 15 P12
    s = add_blank(prs)
    title_block(s, "P12 — Final security validation", "Overall: PASS")
    bullets(
        s,
        [
            "Auth/RBAC · Response safety · API/input · Persistence/recovery",
            "Model-serving · Simulation isolation · Empirical integrity · Audit",
            "Production boundary · Frozen research baseline hashes",
            "Outcome: productionized security-validated prototype — not enterprise SOC claim",
        ],
        size=17,
    )
    footer(s)

    # 16 Demo
    s = add_blank(prs)
    title_block(s, "System demonstration", "See docs/DEMO.md for exact clicks")
    bullets(
        s,
        [
            "UI http://127.0.0.1:5173  ·  API http://127.0.0.1:8000/docs",
            "Flow: Detection → classification / SHAP → risk → incident",
            "Propose → dry-run → approve (CONTROLLED) → verify",
            "Simulation lifecycle · System / ops readiness",
        ],
        size=17,
    )
    footer(s)

    # 17 Limitations
    s = add_blank(prs)
    title_block(s, "Limitations & future work")
    bullets(
        s,
        [
            "CICIDS2017 IID ≠ continuous enterprise traffic · no zero-day claim",
            "Temporal multiclass incomplete for some families · external dataset = future work",
            "Simulation efficacy = assumptions · P11 = CONTROLLED lab, not live EDR",
            "Future: frozen cross-dataset eval · broader temporal families · opt-in sandbox enforcement later",
        ],
        size=16,
    )
    footer(s)

    # 18 Conclusion
    s = add_blank(prs)
    title_block(s, "Conclusion")
    bullets(
        s,
        [
            "Aegis IDS: productionized, security-validated IDS prototype",
            "Frozen ML + XAI + risk + controlled response + simulation + P11/P12 evidence",
            "Hard boundary against live network enforcement preserved",
            "Questions?  docs/VIVA_QA.md  ·  docs/PROJECT_REPORT.md  ·  DEMO.md",
        ],
        top=1.75,
        size=17,
    )
    box = s.shapes.add_textbox(Inches(0.7), Inches(5.5), Inches(12), Inches(1.2))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = (
        "Thank you\n"
        "v1.1-research  ·  v2.0-aegis-productionized  ·  P12 Overall PASS"
    )
    _set_run(r, 16, False, MUTED)
    footer(s, "Fill Team / College / Year on slide 1 before the viva")

    prs.save(OUT)
    print(f"Wrote {OUT} ({OUT.stat().st_size} bytes, {len(prs.slides)} slides)")


if __name__ == "__main__":
    main()
